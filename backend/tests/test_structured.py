import json

import httpx
import pytest

from rag_platform.catalog import query_terms
from rag_platform.config import settings
from rag_platform.evaluate import evaluate
from rag_platform.parsing import content_hash, delete_parse_cache, parse_document, parser_configuration
from rag_platform.service import RagService
from rag_platform.structured import ingest_structured


@pytest.fixture
def source_root(tmp_path, monkeypatch):
    monkeypatch.setattr(settings(), "storage_root", tmp_path)
    return tmp_path


def test_structured_markdown_table_headings_and_cached_retry(source_root, monkeypatch):
    (source_root / "policy.md").write_text(
        "# Perjalanan Dinas\n\n## Akomodasi Domestik\n\n"
        "| Golongan | Plafon Hotel |\n|---|---|\n| Staf | Rp750.000 |\n| Manajer | Rp950.000 |\n\n"
        "## Transportasi\n\nTaksi diganti sesuai kuitansi.")
    configuration = parser_configuration()
    result = parse_document("policy.md", configuration)
    assert not result["cache_hit"]
    hotel = next(c for c in result["chunks"] if "Rp750.000" in c)
    assert "Plafon Hotel" in hotel and "Staf" in hotel
    assert "Akomodasi Domestik" in result["chunk_metadata"][content_hash(hotel)]["headings"]
    assert "Transportasi" in result["catalog_text"]

    def unexpected_conversion(*args):
        raise AssertionError("Retry must reuse parsed source")

    monkeypatch.setattr("rag_platform.parsing.convert_document", unexpected_conversion)
    cached = parse_document("policy.md", configuration)
    assert cached["cache_hit"] and cached["chunks"] == result["chunks"]
    with pytest.raises(AssertionError):
        parse_document("policy.md", {**configuration, "chunk_size": 900})
    delete_parse_cache("policy.md")
    assert not list(source_root.glob("*.parsed.json"))
    assert (source_root / "policy.md").exists()


def test_cache_invalidated_when_original_changes(source_root):
    path = source_root / "source.txt"
    path.write_text("Original travel policy.")
    first = parse_document("source.txt", parser_configuration())
    path.write_text("Updated travel policy.")
    second = parse_document("source.txt", parser_configuration())
    assert not second["cache_hit"] and first["chunks"] != second["chunks"]


def test_token_like_document_text_is_literal_and_parser_upgrade_requires_reindex(source_root):
    (source_root / "tokens.txt").write_text("The example contains <|endoftext|> as literal text.")
    config = parser_configuration()
    assert "<|endoftext|>" in parse_document("tokens.txt", config)["chunks"][0]
    with pytest.raises(ValueError, match="reindex"):
        parse_document("tokens.txt", {**config, "docling_version": "old"})


def test_docx_structure(source_root):
    from docx import Document
    document = Document()
    document.add_heading("Travel", 1)
    document.add_heading("Hotel", 2)
    document.add_paragraph("The domestic allowance is Rp750.000.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Role", "Limit"
    table.cell(1, 0).text, table.cell(1, 1).text = "Staff", "Rp750.000"
    document.save(source_root / "policy.docx")
    result = parse_document("policy.docx", parser_configuration())
    assert any("Hotel" in m["headings"] for m in result["chunk_metadata"].values())
    assert any("Staff" in c and "Limit" in c for c in result["chunks"])


def test_large_section_is_split_with_heading_context(source_root):
    (source_root / "large.md").write_text("# Policy\n\n## Accommodation\n\n" + " ".join(f"Hotel allowance item {i} applies per night." for i in range(300)))
    config = {**parser_configuration(), "chunk_size": 128}
    result = parse_document("large.md", config)
    import tiktoken
    encoder = tiktoken.get_encoding("cl100k_base")
    assert len(result["chunks"]) > 1
    assert all(len(encoder.encode(c)) <= 128 for c in result["chunks"])
    assert all("Accommodation" in m["headings"] for m in result["chunk_metadata"].values())


def test_query_terms_remove_syntax_and_question_words():
    assert query_terms("Berapa plafon hotel untuk perjalanan dinas domestik?") == ["plafon", "hotel", "perjalanan", "dinas", "domestik"]
    assert query_terms("'hotel':* | ! hotel; DROP --") == ["hotel", "drop"]
    assert query_terms("??? berapa yang") == []


@pytest.mark.asyncio
async def test_adapter_preserves_chunks_and_source_identity():
    from types import SimpleNamespace

    import tiktoken

    class Rag:
        tokenizer = SimpleNamespace(encode=tiktoken.get_encoding("cl100k_base").encode)

        async def apipeline_enqueue_documents(self, text, **kwargs):
            self.enqueued = kwargs
            self.chunks = self.chunking_func(self.tokenizer, text, None, False, 0, 1000)

        async def apipeline_process_enqueue_documents(self):
            self.processed = True

    rag = Rag()
    chunks = ["Policy\nHotel\nDomestic limit Rp750.000.", "Policy\nFlights\nEconomy class."]
    await ingest_structured(rag, {"text": "Full document", "chunks": chunks, "track_id": "job",
                                 "configuration": {"parser": {"chunk_size": 1000}}}, "document", "document:version")
    assert rag.enqueued == {"ids": ["document"], "file_paths": ["document:version"], "track_id": "job", "process_options": "C"}
    assert [c["content"] for c in rag.chunks] == chunks and rag.processed


@pytest.mark.asyncio
async def test_retrieval_preserves_provenance_and_rejects_wrong_version(monkeypatch):
    content = "Hotel allowance is Rp750.000."
    snapshot = {"document_id": "doc", "version_id": "v2", "title": "Policy",
                "chunk_metadata": {content_hash(content): {"pages": [3], "headings": ["Hotel"]}}}

    async def operate(self, payload):
        assert "chunk_metadata" not in payload
        return {"data": {"chunks": [{"content": content, "chunk_id": "chunk", "file_path": "doc:v2"},
                                     {"content": "Revoked evidence", "file_path": "doc:v1"}]}}

    monkeypatch.setattr(RagService, "operate", operate)
    sources = await RagService().retrieve("Hotel?", [snapshot], {})
    assert len(sources) == 1 and sources[0]["pages"] == [3] and sources[0]["headings"] == ["Hotel"]


def test_evaluation_measures_routing_and_detects_leakage():
    payload = {"answer": "Rp750.000 [S1]", "abstained": False, "citations": [{"source_id": "S1", "quote": "Rp750.000"}],
               "sources": [{"document_id": "allowed", "headings": ["Hotel"]}],
               "routing": {"document_ids": ["private", "allowed"]}, "timings": {"routing_ms": 2}}
    with httpx.Client(base_url="http://test", transport=httpx.MockTransport(lambda r: httpx.Response(200, json=payload))) as client:
        cases, metrics = evaluate(client, [{"id": "router", "question": "Hotel?", "expected_document_ids": ["allowed"],
                                           "allowed_source_ids": ["allowed"], "forbidden_document_ids": ["private"], "require_provenance": True}])
    assert not cases[0]["passed"] and not cases[0]["checks"]["routing_excludes_forbidden"]
    assert metrics["routing_recall"] == 1 and metrics["routing_precision"] == 0.5 and metrics["routing_mrr"] == 0.5
    assert metrics["p95_routing_ms"] == 2
    assert json.dumps(metrics)
