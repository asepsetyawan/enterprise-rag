import hashlib
import os
import re
import tempfile
import uuid
import zipfile
from pathlib import Path

from fastapi import HTTPException, UploadFile

from .config import settings

ALLOWED = {".txt", ".md", ".pdf", ".docx"}


def file_path(relative: str) -> Path:
    root = settings().storage_root.resolve()
    candidate = (root / relative).resolve()
    if not candidate.is_relative_to(root):
        raise ValueError("Storage path outside configured root")
    return candidate


async def save_upload(upload: UploadFile, tenant: uuid.UUID, document: uuid.UUID, version: uuid.UUID):
    filename = Path(upload.filename or "document.txt").name
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED or len(filename) > 240 or any(ord(c) < 32 for c in filename):
        raise HTTPException(415, "Supported files: TXT, Markdown, text PDF, DOCX")
    relative = f"{tenant.hex}/{document.hex}/{version.hex}{extension}"
    target = file_path(relative)
    target.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    checksum = hashlib.sha256()
    total = 0
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=target.parent, delete=False) as stream:
            temporary = Path(stream.name)
            os.chmod(temporary, 0o600)
            while chunk := await upload.read(1024 * 1024):
                total += len(chunk)
                if total > settings().max_upload_bytes:
                    raise HTTPException(413, "File exceeds upload limit")
                checksum.update(chunk)
                stream.write(chunk)
            stream.flush()
            os.fsync(stream.fileno())
        if not total:
            raise HTTPException(400, "Empty document")
        if extension == ".pdf" and temporary.read_bytes()[:5] != b"%PDF-":
            raise HTTPException(415, "Invalid PDF signature")
        if extension == ".docx":
            try:
                with zipfile.ZipFile(temporary) as archive:
                    entries = archive.infolist()
                    expanded = sum(entry.file_size for entry in entries)
                    if len(entries) > 5000 or expanded > 100 * 1024 * 1024 or expanded > max(total * 100, 2 * 1024 * 1024):
                        raise HTTPException(413, "DOCX expansion exceeds safety limits")
                    if "word/document.xml" not in archive.namelist():
                        raise HTTPException(415, "Not a Word document")
            except zipfile.BadZipFile as exc:
                raise HTTPException(415, "Invalid DOCX archive") from exc
        os.replace(temporary, target)
        directory = os.open(target.parent, os.O_RDONLY)
        try:
            os.fsync(directory)
        finally:
            os.close(directory)
        return filename, relative, checksum.hexdigest(), total
    finally:
        if temporary and temporary.exists():
            temporary.unlink()
        await upload.close()


def extract_text(relative: str) -> str:
    path = file_path(relative)
    extension = path.suffix.lower()
    if extension in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8-sig")
    elif extension == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted or len(reader.pages) > 1000:
            raise ValueError("Encrypted or oversized PDF is not supported")
        parts = []
        length = 0
        for page in reader.pages:
            part = page.extract_text() or ""
            length += len(part)
            if length > settings().max_text_chars:
                raise ValueError("Extracted text exceeds limit")
            parts.append(part)
        text = "\n\n".join(parts)
    elif extension == ".docx":
        from docx import Document
        document = Document(path)
        parts = [p.text for p in document.paragraphs]
        parts.extend(" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
        text = "\n\n".join(parts)
    else:
        raise ValueError("Unsupported file format")
    text = text.replace("\x00", "").strip()
    if not text or len(text) > settings().max_text_chars:
        raise ValueError("No usable text or extracted text exceeds limit; scanned files need an OCR pipeline")
    return text


def workspace_for(tenant: uuid.UUID, document: uuid.UUID, build: uuid.UUID):
    # Distinct per-attempt immutable namespace: stale workers can never overwrite a published index.
    return f"t{tenant.hex}_d{document.hex}_b{build.hex}"


def validate_workspace(value: str, tenant: uuid.UUID, document: uuid.UUID):
    expected = f"t{tenant.hex}_d{document.hex}_b"
    if not value.startswith(expected) or not re.fullmatch(r"t[0-9a-f]{32}_d[0-9a-f]{32}_b[0-9a-f]{32}", value):
        raise ValueError("Workspace does not belong to tenant/document")
